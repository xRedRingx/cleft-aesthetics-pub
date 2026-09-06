"""Load the rater score sheet.

Phase 1 brief §1.3. Two workbooks exist, identical in content: one spells the
grades out, one stores integers. This loader accepts either and produces the same
result.

    Excellent = 1   Good = 2   Fair = 3   Poor = 4   Very poor = 5

Verified across all 1,255 cells (251 rows x 5 raters) with zero missing.

Two details make exact string matching fail on the real file while passing on a
tidy fixture, so both are handled explicitly: some rater headers contain runs of
whitespace, and "Very poor" has a lowercase p.

The sheet also carries ``Average`` and ``Median`` columns. They are **not read**.
Every derived quantity is recomputed from the raw grades, because a precomputed
column that has drifted from the cells beside it disagrees silently.

**[AMENDED 2026-08-17]** ``load_median`` now reads the ``Median`` column --
for ONE consumer, Phase 10's CleftGNN replication, whose training label is
that column by the group's own design (``phase10.CONSENSUS_LABEL_IS_THE_MEDIAN``).
The paragraph above still governs: the reader recomputes the median from the
five rater cells and **refuses any row where the two disagree**, so the
precomputed column is verified rather than trusted. ``load`` is unchanged and
still ignores it.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np

#: Canonical rater column names, in the order the grade matrix uses.
RATERS: tuple[str, ...] = (
    "Rater 7 - Cleft patient",
    "Rater 8 - Orthodontist",
    "Rater 9 - Speech and language therapist",
    "Rater 10 - Plastic surgeon",
    "Rater 11 - Psychologist",
)

ID_COLUMN = "RanaPhotoID"

#: Ignored on purpose by ``load`` -- see the module docstring. ``Median`` is
#: read by ``load_median`` alone, and only against its own recomputation.
DERIVED_COLUMNS = ("Average", "Median")

MEDIAN_COLUMN = "Median"

GRADE_WORDS: dict[str, int] = {
    "excellent": 1,
    "good": 2,
    "fair": 3,
    "poor": 4,
    "very poor": 5,
}

GRADE_MIN, GRADE_MAX = 1, 5


class ScoreSheetError(ValueError):
    """The score sheet is not what it must be. Always fatal."""


def normalise_header(text: object) -> str:
    """Collapse whitespace runs and strip, so header quirks stop mattering."""
    return re.sub(r"\s+", " ", str(text if text is not None else "")).strip()


def parse_grade(value: object) -> int:
    """One cell to an integer grade. Accepts a word or a number, nothing else."""
    if value is None:
        raise ScoreSheetError("missing grade cell")

    if isinstance(value, bool):
        raise ScoreSheetError(f"grade is a boolean: {value!r}")

    if isinstance(value, (int, float)):
        if float(value) != int(value):
            raise ScoreSheetError(f"grade is not a whole number: {value!r}")
        grade = int(value)
    else:
        text = normalise_header(value)
        if not text:
            raise ScoreSheetError("missing grade cell")
        word = text.casefold()
        if word in GRADE_WORDS:
            return GRADE_WORDS[word]
        if re.fullmatch(r"-?\d+", text):
            grade = int(text)
        else:
            raise ScoreSheetError(
                f"unrecognised grade {value!r}. Expected one of "
                f"{sorted(GRADE_WORDS)} (any capitalisation) or an integer "
                f"{GRADE_MIN}-{GRADE_MAX}."
            )

    if not GRADE_MIN <= grade <= GRADE_MAX:
        raise ScoreSheetError(
            f"grade {grade} is outside the {GRADE_MIN}-{GRADE_MAX} range"
        )
    return grade


@dataclass(frozen=True)
class ScoreRow:
    photo_id: int
    grades: tuple[int, ...]


@dataclass
class ScoreSheet:
    rows: dict[int, ScoreRow]
    raters: tuple[str, ...]
    n_missing: int = 0

    @property
    def n_cells(self) -> int:
        return len(self.rows) * len(self.raters)

    def photo_ids(self) -> list[int]:
        return sorted(self.rows)

    def grade_matrix(self) -> np.ndarray:
        """(n_rows, n_raters) integer grades, rows ordered by photo id."""
        return np.array(
            [self.rows[pid].grades for pid in self.photo_ids()], dtype=np.int64
        )


def _read_workbook(path: Path) -> list[list[object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _locate_columns(header: list[object]) -> tuple[int, dict[str, int]]:
    normalised = [normalise_header(cell) for cell in header]
    lookup = {name.casefold(): index for index, name in enumerate(normalised)}

    if ID_COLUMN.casefold() not in lookup:
        raise ScoreSheetError(
            f"no {ID_COLUMN!r} column. Found: {[n for n in normalised if n]}"
        )

    rater_index: dict[str, int] = {}
    missing = []
    for rater in RATERS:
        index = lookup.get(rater.casefold())
        if index is None:
            missing.append(rater)
        else:
            rater_index[rater] = index

    if missing:
        raise ScoreSheetError(
            f"missing rater column(s): {missing}. Found: "
            f"{[n for n in normalised if n]}. Headers are normalised for "
            "whitespace before matching, so this is a genuine absence."
        )

    return lookup[ID_COLUMN.casefold()], rater_index


def load(path: str | Path, *, expected_rows: int | None = None) -> ScoreSheet:
    """Load and validate a score-sheet workbook.

    ``expected_rows`` asserts the measured row count (251 for the real sheet).
    """
    path = Path(path)
    table = _read_workbook(path)
    if not table:
        raise ScoreSheetError(f"{path} is empty")

    id_index, rater_index = _locate_columns(table[0])

    rows: dict[int, ScoreRow] = {}
    for line_number, raw in enumerate(table[1:], start=2):
        if all(cell is None for cell in raw):
            continue

        photo_id = _parse_photo_id(raw, id_index, line_number, path)
        if photo_id in rows:
            raise ScoreSheetError(
                f"{path} row {line_number}: duplicate {ID_COLUMN} {photo_id}. "
                "Each photograph is scored exactly once."
            )

        grades = []
        for rater in RATERS:
            cell = raw[rater_index[rater]] if rater_index[rater] < len(raw) else None
            try:
                grades.append(parse_grade(cell))
            except ScoreSheetError as exc:
                raise ScoreSheetError(
                    f"{path} row {line_number} ({ID_COLUMN} {photo_id}), "
                    f"column {rater!r}: {exc}"
                ) from None

        rows[photo_id] = ScoreRow(photo_id=photo_id, grades=tuple(grades))

    if not rows:
        raise ScoreSheetError(f"{path} has a header but no data rows")

    if expected_rows is not None and len(rows) != expected_rows:
        raise ScoreSheetError(
            f"{path}: expected {expected_rows} scored rows, found {len(rows)}"
        )

    return ScoreSheet(rows=rows, raters=RATERS, n_missing=0)


def _parse_photo_id(raw: list[object], index: int, line_number: int, path: Path) -> int:
    value = raw[index] if index < len(raw) else None
    if value is None or normalise_header(value) == "":
        raise ScoreSheetError(f"{path} row {line_number}: missing {ID_COLUMN}")
    text = normalise_header(value)
    if not re.fullmatch(r"\d+", text):
        raise ScoreSheetError(
            f"{path} row {line_number}: {ID_COLUMN} {value!r} is not an integer"
        )
    return int(text)


def load_median(path: str | Path, *, expected_rows: int | None = None) -> dict[int, int]:
    """``{photo_id: Median}`` -- the sheet's own column, VERIFIED cell by cell.

    Phase 10's CleftGNN replication trains on the group's consensus grade,
    which is this column (``phase10.CONSENSUS_LABEL_IS_THE_MEDIAN``: the
    derived files' "consensus" column IS the sheet's Median, verified by the
    operator on rows 241/243/245 against APScores' eighth column). It is read
    HERE rather than by a second ad-hoc reader, because this module already
    knows the sheet's identity -- its id column, its five rater columns, and
    the whitespace and capitalisation quirks that make exact matching fail on
    the real file.

    **The precomputed column is verified, never trusted** -- the module
    docstring's rule, kept: the median of the five rater cells is recomputed
    for every row and any disagreement RAISES. Two internal stories about the
    same row is a fault to stop on, not a value to pick between. The median of
    five integers is the third of the sorted five, so no tie rule exists to
    get wrong.
    """
    path = Path(path)
    table = _read_workbook(path)
    if not table:
        raise ScoreSheetError(f"{path} is empty")

    id_index, rater_index = _locate_columns(table[0])
    normalised = [normalise_header(cell) for cell in table[0]]
    matches = [
        index for index, name in enumerate(normalised)
        if name.casefold() == MEDIAN_COLUMN.casefold()
    ]
    if len(matches) != 1:
        raise ScoreSheetError(
            f"{path}: expected exactly one {MEDIAN_COLUMN!r} column, found "
            f"{len(matches)}. Header: {[n for n in normalised if n]}"
        )
    median_index = matches[0]

    out: dict[int, int] = {}
    for line_number, raw in enumerate(table[1:], start=2):
        if all(cell is None for cell in raw):
            continue
        photo_id = _parse_photo_id(raw, id_index, line_number, path)
        if photo_id in out:
            raise ScoreSheetError(
                f"{path} row {line_number}: duplicate {ID_COLUMN} {photo_id}"
            )

        grades = []
        for rater in RATERS:
            cell = raw[rater_index[rater]] if rater_index[rater] < len(raw) else None
            try:
                grades.append(parse_grade(cell))
            except ScoreSheetError as exc:
                raise ScoreSheetError(
                    f"{path} row {line_number} ({ID_COLUMN} {photo_id}), "
                    f"column {rater!r}: {exc}"
                ) from None
        recomputed = sorted(grades)[len(grades) // 2]

        cell = raw[median_index] if median_index < len(raw) else None
        try:
            published = parse_grade(cell)
        except ScoreSheetError as exc:
            raise ScoreSheetError(
                f"{path} row {line_number} ({ID_COLUMN} {photo_id}), "
                f"column {MEDIAN_COLUMN!r}: {exc}"
            ) from None

        if published != recomputed:
            raise ScoreSheetError(
                f"{path} row {line_number} ({ID_COLUMN} {photo_id}): the "
                f"{MEDIAN_COLUMN!r} column says {published} but the five "
                f"rater cells {grades} give {recomputed}. The sheet tells two "
                "stories about one row; nothing here picks between them."
            )
        out[photo_id] = published

    if not out:
        raise ScoreSheetError(f"{path} has a header but no data rows")
    if expected_rows is not None and len(out) != expected_rows:
        raise ScoreSheetError(
            f"{path}: expected {expected_rows} scored rows, found {len(out)}"
        )
    return out
